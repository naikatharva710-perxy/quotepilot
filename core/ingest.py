"""Turn whatever a vendor sent into something the model can read, while keeping
addresses (cell refs, paragraph numbers, pages, line numbers) so every extracted
value can point back to where it came from."""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path

MIME = {
    ".pdf": "application/pdf",
    ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".webp": "image/webp",
}


@dataclass
class Doc:
    name: str
    kind: str                     # excel | pdf | word | image | email | text
    text: str = ""                # addressed text representation (may be empty for images)
    binary: bytes | None = None   # sent natively to the model (pdf / image)
    mime: str | None = None
    notes: list[str] = field(default_factory=list)


def _excel(name: str, data: bytes) -> Doc:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"=== Sheet '{ws.title}' ===")
        for row in ws.iter_rows():
            cells = [f"{c.coordinate}={c.value}" for c in row if c.value not in (None, "")]
            if cells:
                out.append(" | ".join(cells))
    return Doc(name, "excel", "\n".join(out))


def _word(name: str, data: bytes) -> Doc:
    import docx
    d = docx.Document(io.BytesIO(data))
    out, n = [], 0
    for p in d.paragraphs:
        if p.text.strip():
            n += 1
            out.append(f"[para {n}] {p.text.strip()}")
    for ti, t in enumerate(d.tables, 1):
        for ri, row in enumerate(t.rows, 1):
            out.append(f"[table {ti} row {ri}] " + " | ".join(c.text.strip() for c in row.cells))
    return Doc(name, "word", "\n".join(out))


def _pdf(name: str, data: bytes) -> Doc:
    import pdfplumber
    out = []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for i, page in enumerate(pdf.pages, 1):
                txt = page.extract_text() or ""
                out.append(f"=== page {i} ===\n{txt}")
    except Exception as e:  # scanned / odd PDFs still go to the model natively
        out.append(f"(text layer unreadable: {e})")
    doc = Doc(name, "pdf", "\n".join(out), data, "application/pdf")
    if sum(len(x) for x in out) < 200:
        doc.notes.append("PDF has little or no text layer; relying on the model's visual read")
    return doc


def _text(name: str, data: bytes) -> Doc:
    raw = data.decode("utf-8", errors="replace")
    lines = [f"[line {i}] {l}" for i, l in enumerate(raw.splitlines(), 1) if l.strip()]
    kind = "email" if raw.lower().startswith(("from:", "subject:", "to:")) or name.endswith(".eml") else "text"
    return Doc(name, kind, "\n".join(lines))


def load(name: str, data: bytes) -> Doc:
    ext = Path(name).suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        return _excel(name, data)
    if ext == ".docx":
        return _word(name, data)
    if ext == ".pdf":
        return _pdf(name, data)
    if ext in (".jpg", ".jpeg", ".png", ".webp"):
        return Doc(name, "image", "", data, MIME[ext], ["Photo: values are read visually; expect lower confidence"])
    if ext in (".txt", ".eml", ".csv", ".md"):
        return _text(name, data)
    raise ValueError(f"Unsupported file type: {ext}")
